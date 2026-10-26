#!/usr/bin/env python3
"""Sixth check axis: is the text in a workbook actually VISIBLE?

The other five axes all read values. `verify_workbook.py` compares figures,
`check_builds_current.py` compares timestamps, `audit_gaps.py` counts filled
cells, `check_ratio_ties.py` and `check_overview_ties.py` divide and compare
them. None of them looks at presentation, so a workbook can be correct,
current, complete and internally consistent while delivering its evidence
invisibly. That is GA-030: 1,743 merged note and source cells across 140 of
145 workbooks held text that did not fit the explicit row height set under
it, and an explicit height disables Excel's auto-fit, so the overflow is
simply not drawn.

What it checks: every merged, wrapped, single-row cell against the height its
text needs at that merged width.

The capacity model is NOT reimplemented here. It is imported from
`bank_workbook._fit_height` - the same function the builders size rows with -
so this tool cannot drift from the thing it audits. Its two constants were
measured in Excel rather than modelled; see that function.

    python3 scripts/check_cell_fit.py                 # whole corpus
    python3 scripts/check_cell_fit.py --bank MONUMENT # one workbook
    python3 scripts/check_cell_fit.py --tsv           # one row per clipped cell

Like the other axes this reports LEADS, NOT ERRORS. The model is an estimate
of Excel's layout, not Excel's layout: a cell flagged at 1.05x may well be
fine. The tolerance below exists for that reason. A cell at 20x is clipped
under any model.
"""

import argparse
import glob
import os
import sys

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from bank_workbook import BankWorkbook  # noqa: E402

# Flag only cells needing more than this multiple of the height they have, so
# that estimation error around the boundary does not produce noise.
TOLERANCE = 1.2
# Excel's own default when a row carries no explicit height.
DEFAULT_ROW_HEIGHT = 15.0


def clipped_cells(path):
    """Yield (sheet, row, chars, have_pt, need_pt) for each clipped cell."""
    wb = openpyxl.load_workbook(path)
    try:
        for ws in wb.worksheets:
            widths = {}
            for letter, dim in ws.column_dimensions.items():
                if dim.width:
                    widths[letter] = dim.width
            for rng in ws.merged_cells.ranges:
                if rng.min_row != rng.max_row:
                    continue
                cell = ws.cell(row=rng.min_row, column=rng.min_col)
                if not cell.value or not isinstance(cell.value, str):
                    continue
                if not (cell.alignment and cell.alignment.wrap_text):
                    continue
                width_units = 0.0
                for col in range(rng.min_col, rng.max_col + 1):
                    letter = openpyxl.utils.get_column_letter(col)
                    width_units += widths.get(letter, 8.43)
                need = BankWorkbook._fit_height(cell.value, width_units)
                if need is None:
                    continue
                dim = ws.row_dimensions.get(rng.min_row)
                have = (
                    dim.height
                    if dim is not None and dim.height
                    else DEFAULT_ROW_HEIGHT
                )
                if need > have * TOLERANCE:
                    yield ws.title, rng.min_row, len(cell.value), have, need
    finally:
        wb.close()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--bank", help="substring of the workbook filename")
    ap.add_argument("--tsv", action="store_true", help="one row per clipped cell")
    args = ap.parse_args()

    paths = sorted(glob.glob("banks/*FINANCIALS.xlsx"))
    if args.bank:
        needle = args.bank.upper()
        paths = [p for p in paths if needle in os.path.basename(p).upper()]
        if not paths:
            print(f"No workbook matches {args.bank!r}", file=sys.stderr)
            return 2

    if args.tsv:
        print("bank\tsheet\trow\tchars\thave_pt\tneed_pt\tratio")

    total_cells = 0
    total_hidden = 0
    affected = []
    for path in paths:
        bank = os.path.basename(path).replace(" FINANCIALS.xlsx", "")
        rows = list(clipped_cells(path))
        if not rows:
            continue
        affected.append((bank, len(rows)))
        for sheet, row, chars, have, need in rows:
            total_cells += 1
            # characters that do not fit, by the same proportion
            total_hidden += int(chars * max(0.0, 1 - have / need))
            if args.tsv:
                print(
                    f"{bank}\t{sheet}\t{row}\t{chars}\t{have:.0f}\t"
                    f"{need:.0f}\t{need / have:.1f}"
                )

    if not args.tsv:
        for bank, n in sorted(affected, key=lambda t: -t[1]):
            print(f"{n:5d}  {bank}")
        print()
        print(
            f"{total_cells} clipped cells across {len(affected)} of "
            f"{len(paths)} workbooks; ~{total_hidden:,} characters not visible "
            f"without a manual row resize."
        )
    return 1 if total_cells else 0


if __name__ == "__main__":
    sys.exit(main())
