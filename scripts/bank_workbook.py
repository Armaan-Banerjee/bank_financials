"""
Shared helpers for building "<BANK> FINANCIALS.xlsx" workbooks in banks/.

Standard structure (as of the ST- wayfinder map, see
wayfinder/statements/map.md - 18 sheets total): an "Overview" sheet
(headline Balance Sheet/P&L/Equity/Cash Flow figures plus Pillar 3 ratios -
see the "spend and risk" lens in the map's Destination), then "Balance
Sheet" (Consolidated Statement of Financial Position), "Profit & Loss"
(Consolidated Statement of Comprehensive Income), a "Statement of Changes
in Equity" sheet, a "Cash Flow Statement" sheet (full statement, period
columns, most-recent-first), an "Asset Quality" sheet (loan book by
product and IFRS 9 stage, plus NPL/coverage ratios), one sheet per Pillar 3
key metric (CET1 Capital, CET1 Ratio, Tier 1 Capital, Tier 1 Ratio, Total
Capital, Total Capital Ratio, Total RWAs, Leverage Ratio, LCR, NSFR, MREL
Ratio), and finally an "RWA Breakdown" sheet (placed right after Total
RWAs, not after MREL Ratio, since it's itself a Pillar 3 disclosure). Each
sheet ends in a single merged, wrapped source-citation cell rather than a
per-row citation column.

Usage:
    from bank_workbook import BankWorkbook, PILLAR3_SHEET_NAMES

    wb = BankWorkbook(
        bank_name="Example Bank Plc",
        years=["FY2025", "FY2024", "FY2023"],
        header_color="0A2540",
    )
    wb.add_overview_sheet(balance_sheet_totals=[...], balance_sheet_unit=..., income_statement_totals=[...],
                           income_statement_unit=..., equity_changes_totals=[...], equity_changes_unit=...,
                           cash_flow_totals=[...], cash_flow_unit=..., ratios=[...])
    wb.add_balance_sheet_sheet(title=..., subtitle=..., rows=[...], sources_text=...)
    wb.add_income_statement_sheet(title=..., subtitle=..., rows=[...], sources_text=...)
    wb.add_equity_changes_sheet(title=..., subtitle=..., headers=[...], rows=[...], sources_text=...)
    wb.add_cash_flow_sheet(title=..., subtitle=..., rows=[...], sources_text=...)
    wb.add_asset_quality_sheet(title=..., subtitle=..., rows=[...], sources_text=...)
    wb.add_metric_sheet("CET1 Capital", unit="£'000", rows_data=[...], sources_text=...)
    wb.add_rwa_breakdown_sheet(title=..., subtitle=..., rows=[...], sources_text=...)
    wb.add_long_form_sheet("Interim Pillar 3", headers=[...], rows=[...])
    ...
    wb.save("/Users/armaan/code/katalysis/banks/EXAMPLE BANK FINANCIALS.xlsx")

Row tuple formats:
    Cash flow / Balance Sheet / Profit & Loss / Asset Quality / RWA
        Breakdown `rows`: (kind, label, values) where kind is "SECTION"
        (label-only divider), "DATA" (a normal line item) or "TOTAL"
        (bolded subtotal/total). `values` is a dict of {year: number};
        omit a year to leave it blank. Asset Quality and RWA Breakdown
        ratio rows may use string values (e.g. "13.55%") the same way
        Metric sheet rows do.
    Equity changes `rows`: (kind, label, values) where kind is "DATA" or
        "TOTAL" (bolded - use for balance b/f-c/f and comprehensive-income
        rows). `values` is a tuple positionally aligned to the sheet's
        `headers` (equity components), not a year dict - this sheet reads
        chronologically oldest-to-newest, not year-columned.
    Metric sheet `rows_data`: (label, values) pairs, values as above. Use
        string values (e.g. "82.94%", "Not publicly disclosed") freely -
        cells are written as-is.
"""

import re

import openpyxl
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.chart.layout import Layout, ManualLayout
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

# The bank's own UK KM1 "Key metrics" template. Kept OUT of
# PILLAR3_SHEET_NAMES for the same reason "RWA Breakdown" is: that list names
# the 11 single-metric sheets, whereas this is a published disclosure template
# reproduced whole. It sits immediately before those 11 sheets - sheet order
# follows call order in the build script, so call add_km1_sheet() before the
# first add_metric_sheet().
KM1_SHEET_NAME = "KM1 Key Metrics"

PILLAR3_SHEET_NAMES = [
    "CET1 Capital",
    "CET1 Ratio",
    "Tier 1 Capital",
    "Tier 1 Ratio",
    "Total Capital",
    "Total Capital Ratio",
    "Total RWAs",
    "Leverage Ratio",
    "LCR",
    "NSFR",
    "MREL Ratio",
]

_PERCENT_RE = re.compile(r"^\s*(-?[\d,]+\.?\d*)\s*%")


def _parse_percent(value):
    """Parse a displayed ratio value ("14.5%") into a float (14.5) for
    charting; returns None for non-numeric values ("Not disclosed", etc.)."""
    if isinstance(value, (int, float)):
        return value
    if isinstance(value, str):
        m = _PERCENT_RE.match(value)
        if m:
            return float(m.group(1).replace(",", ""))
    return None


THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
SECTION_FONT = Font(bold=True)
TOTAL_FONT = Font(bold=True)
TITLE_FONT = Font(bold=True, size=13)
SUBTITLE_FONT = Font(italic=True, size=9, color="666666")
SOURCE_FONT = Font(italic=True, size=9, color="444444")


class BankWorkbook:
    def __init__(self, bank_name, years, year_label=None, header_color="0A2540"):
        """
        bank_name: used as the "<Bank Name> — <Sheet>" title on metric sheets.
        years: list of year keys (e.g. "FY2025"), most recent first - this
            order drives column order on every sheet.
        year_label: optional {year_key: column_header_text} override, e.g. to
            append entity/period flags like "FY2024*". Defaults to the key.
        header_color: hex fill color (no '#') for the header row - vary this
            per bank so workbooks are visually distinguishable.
        """
        self.bank_name = bank_name
        self.years = years
        self.year_label = year_label or {y: y for y in years}
        self.header_fill = PatternFill(
            start_color=header_color, end_color=header_color, fill_type="solid"
        )
        self.header_font = Font(color="FFFFFF", bold=True)
        self.wb = openpyxl.Workbook()
        self._first_sheet_used = False

    # -- internal styling helpers -------------------------------------------------
    def _style_header(self, ws, row, ncols):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=row, column=c)
            cell.fill = self.header_fill
            cell.font = self.header_font
            cell.alignment = Alignment(
                horizontal="left", vertical="center", wrap_text=True
            )

    def _autosize(self, ws, widths):
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w

    def _write_source_cell(self, ws, row, ncols, text, height=150):
        # openpyxl silently truncates any cell string past 32,767 chars
        # (its own hard limit) rather than raising - caught the hard way in
        # HD-078 (Co-operative Bank), where citation text had grown past it
        # unnoticed. Fail loudly instead.
        if text and len(text) > 32000:
            raise ValueError(
                f"Source citation text for sheet {ws.title!r} is {len(text)} "
                "chars, over openpyxl's 32,767-char cell limit (32,000 "
                "warning threshold) - it would be silently truncated. Trim "
                "it before saving."
            )
        cell = ws.cell(row=row, column=1, value=text)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
        cell.font = SOURCE_FONT
        cell.alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")
        ws.row_dimensions[row].height = height

    def append_source_cell(self, ws, text, ncols=None, height=150):
        """Add a FURTHER merged source-citation cell below the one a sheet
        already has.

        Every sheet ends in a single merged citation cell, and openpyxl caps a
        cell at 32,767 characters - a ceiling the Co-operative Bank's shared
        statement-source block had already grown to within a few hundred
        characters of by 2026-09-19, which left no room to cite the 31 newly
        filled Profit & Loss years. Splitting the citation across two adjacent
        cells keeps every source on the sheet that uses it rather than pushing
        provenance out into a file the reader of the workbook never sees.

        Use this only when one cell genuinely cannot hold the citation; the
        one-cell form stays the default.
        """
        if ncols is None:
            ncols = ws.max_column
        self._write_source_cell(ws, ws.max_row + 2, ncols, text, height=height)
        return ws

    def _style_chart_axes(self, chart):
        # openpyxl's TextAxis/NumericAxis both default axPos to "l", which
        # puts the category axis on top of the value axis - fix the category
        # (x) axis onto the bottom and make sure both axes actually draw
        # their tick labels and lines.
        chart.x_axis.axPos = "b"
        chart.y_axis.axPos = "l"
        chart.x_axis.tickLblPos = "nextTo"
        chart.y_axis.tickLblPos = "nextTo"
        chart.x_axis.delete = False
        chart.y_axis.delete = False
        chart.x_axis.majorTickMark = "out"
        chart.y_axis.majorTickMark = "out"

    def _reserve_bottom_legend(self, chart, legend_h, plot_h, plot_top=0.12):
        # Pin the legend to its own band at the very bottom of the chart
        # area and explicitly shrink the plotting rectangle above it, so
        # Excel's automatic layout can't let a wrapped multi-line legend
        # encroach on the axis tick labels above it.
        chart.legend.layout = Layout(
            manualLayout=ManualLayout(
                xMode="edge",
                yMode="edge",
                x=0.02,
                y=1.0 - legend_h,
                h=legend_h,
                w=0.96,
            )
        )
        # ChartBase._write() overwrites plot_area.layout with chart.layout at
        # save time, so the plot rectangle must be set here, not on plot_area.
        chart.layout = Layout(
            manualLayout=ManualLayout(
                layoutTarget="inner",
                xMode="edge",
                yMode="edge",
                x=0.10,
                y=plot_top,
                w=0.85,
                h=plot_h,
            )
        )

    def _next_sheet(self, name):
        if not self._first_sheet_used:
            ws = self.wb.active
            ws.title = name[:31]
            self._first_sheet_used = True
            return ws
        return self.wb.create_sheet(title=name[:31])

    @staticmethod
    def _has_value(v):
        """True if a cell would print something a reader can read.

        A dash counts. The user's decision of 2026-09-18 is that a printed
        dash stays on the sheet as a literal "-", because it is the bank
        SAYING something ("this row does not apply to us") and is a different
        statement from silence. So a year whose only content is dashes is a
        year the bank published, and must keep its column.
        """
        if v is None:
            return False
        return not (isinstance(v, str) and not v.strip())

    def _trim_trailing_empty_years(self, years, value_dicts):
        """Drop year columns off the OLDEST end that no row on this sheet fills.

        The user's instruction of 2026-09-17 (`bank_probs.txt` general point 1,
        naming Alrayan's KM1, which printed FY2021 and FY2020 headers over two
        entirely empty columns):

            If no Km1 for 2021 and earlier, please dont put header for them in
            spreadsheet

        and the scope decision of 2026-09-18: **trailing only, and on every
        sheet, not only KM1.** The reason a column is empty at the old end is
        the same on every sheet - the disclosure did not exist that far back -
        so the treatment is the same on every sheet.

        TRAILING ONLY IS THE WHOLE POINT, AND THE ASYMMETRY IS DELIBERATE.
        An empty column at the NEWEST end means something completely
        different: the bank has very probably published and we have not
        transcribed it yet. Suppressing that would hide a real gap behind a
        tidy sheet, which is the opposite of what the reader needs - see
        `wayfinder/gaps/map.md`, where the corpus split 662 leading
        sheet-years (real gaps to chase) against 1,881 trailing ones (correct
        as they are). Leading empties keep their headers so they stay visible.

        A sheet where EVERY year is empty is left alone. Such a sheet is all
        leading and all trailing at once, and the honest reading is that it is
        a data gap rather than a presentation problem, so it keeps its columns
        and stays visible. Trimming it would leave a one-column sheet.

        Per-sheet, not per-workbook: year columns may therefore stop at
        different places on different sheets of one workbook. That is intended
        - every column carries its own year label, and no sheet should print a
        header it cannot fill.
        """
        dicts = [d for d in value_dicts if d]
        trimmed = list(years)
        while trimmed:
            if any(self._has_value(d.get(trimmed[-1])) for d in dicts):
                break
            trimmed.pop()
        return trimmed if trimmed else list(years)

    def patch_year_column(self, sheet_name, year, values_by_label, header_row=3):
        """Write one year's values into an ALREADY-BUILT sheet, by row label.

        Three scripts supply a year out-of-band rather than in `rows` - Bank of
        Beirut UK, Gatehouse and National Bank of Egypt UK all carry an FY2017
        block bolted on after the sheet is built, because that year comes from a
        different source than the rest. Each used to compute its target column
        as `1 + YEARS.index(year) + 1`.

        THAT BROKE THE MOMENT COLUMN TRIMMING LANDED, and it broke SILENTLY,
        which is the part worth remembering. `_trim_trailing_empty_years` looks
        at `rows`; a year supplied afterwards is invisible to it, so the column
        was trimmed away and the patch then wrote its values into the cell
        position the year USED to occupy - past the end of the header. The
        figures were still in the file, in a column with no year above them: 98
        cells across the three banks, present, plausible, and unlabelled.

        An index computed from one list and applied to another is only correct
        while nothing can change either list. Resolve the column from the sheet
        as it actually is, and if the year is not there, put it there.
        """
        ws = self.wb[sheet_name]
        headers = [ws.cell(row=header_row, column=c).value
                   for c in range(1, ws.max_column + 1)]
        want = self.year_label.get(year, year)
        col = None
        for i, h in enumerate(headers, start=1):
            # the header may carry a unit suffix - "FY2017 (£'000)" - so match
            # on the label being present rather than on equality
            if h is not None and str(want) in str(h):
                col = i
                break
        if col is None:
            # the year was trimmed (or never built): re-create the column,
            # copying the unit suffix from the neighbouring header so the new
            # one reads the same way as the rest of the row
            col = len(headers) + 1
            suffix = ""
            for h in reversed(headers):
                if h is None:
                    continue
                m = re.search(r"\((.*)\)\s*$", str(h))
                if m:
                    suffix = f" ({m.group(1)})"
                break
            ws.cell(row=header_row, column=col, value=f"{want}{suffix}")
            # reuse the sheet's own header styling rather than re-deriving it,
            # so a restored column is indistinguishable from an original one
            self._style_header(ws, header_row, col)
            ws.cell(row=header_row, column=col).border = BORDER

        labels = {str(ws.cell(row=r, column=1).value).strip(): r
                  for r in range(header_row + 1, ws.max_row + 1)}
        written = 0
        for label, value in values_by_label.items():
            r = labels.get(str(label).strip())
            if r is not None:
                ws.cell(row=r, column=col, value=value)
                ws.cell(row=r, column=col).border = BORDER
                written += 1
        return written

    def _assert_no_orphan_columns(self):
        """Refuse to save a sheet holding data in a column with no header.

        The failure this exists for wrote real figures into unlabelled columns
        and nothing complained: not `verify_workbook.py`, which reads headers
        out to `max_column` and simply reported `None` where a year should be;
        not `check_builds_current.py`, which only compares timestamps; not
        `audit_gaps.py`, which counts cells under headers it can name and so
        never saw them at all. A silent orphan beats every instrument in this
        repo, so the check belongs at the point of writing.
        """
        for ws in self.wb.worksheets:
            if ws.max_row < 4:
                continue
            headers = [ws.cell(row=3, column=c).value
                       for c in range(1, ws.max_column + 1)]
            if not headers or headers[0] is None:
                continue  # not a header-row-3 sheet (Overview, equity roll-forward)
            for c in range(len(headers), 0, -1):
                if headers[c - 1] is not None:
                    break
                orphans = [ws.cell(row=r, column=c).value
                           for r in range(4, ws.max_row + 1)]
                orphans = [v for v in orphans
                           if v is not None and str(v).strip() != ""]
                if orphans:
                    raise ValueError(
                        f"{self.bank_name}: sheet {ws.title!r} holds "
                        f"{len(orphans)} value(s) in column {c}, which has no "
                        f"header. A year column was almost certainly trimmed "
                        f"after something computed this column's index from an "
                        f"untrimmed year list - use patch_year_column() instead "
                        f"of an index into YEARS. First value: {orphans[0]!r}"
                    )

    # -- public API -----------------------------------------------------------
    def _add_statement_sheet(
        self,
        title,
        subtitle,
        rows,
        sources_text,
        sheet_name,
        first_col_width,
        source_height,
        unit_suffix,
        years=None,
    ):
        """Shared body for the full-statement sheets (Cash Flow, Balance
        Sheet, Profit & Loss): a title/subtitle, then a header row of years,
        then one row per (kind, label, values) tuple - kind in {"SECTION",
        "DATA", "TOTAL"} - ending in a merged wrapped source-citation cell.

        years: optional override of self.years for this sheet only - use
            when one sheet's historical depth has been extended further back
            than the rest of the workbook (e.g. a statutory-statement-only
            historical-depth extension that deliberately leaves Pillar 3 /
            Asset Quality / RWA Breakdown at the project-wide floor).
        """
        years = years if years is not None else self.years
        years = self._trim_trailing_empty_years(
            years, (values for _kind, _label, values in rows)
        )
        ws = self._next_sheet(sheet_name)
        ws.freeze_panes = "B4"
        ws["A1"] = title
        ws["A1"].font = TITLE_FONT
        ws["A2"] = subtitle
        ws["A2"].font = SUBTITLE_FONT

        ncols = 1 + len(years)
        headers = ["Line item"] + [
            f"{self.year_label[y]}{unit_suffix}" for y in years
        ]
        header_row = 3
        for c, h in enumerate(headers, start=1):
            ws.cell(row=header_row, column=c, value=h)
        self._style_header(ws, header_row, ncols)

        r = header_row + 1
        for kind, label, values in rows:
            ws.cell(row=r, column=1, value=label)
            if kind == "SECTION":
                ws.cell(row=r, column=1).font = SECTION_FONT
            else:
                for ci, y in enumerate(years, start=2):
                    ws.cell(row=r, column=ci, value=values.get(y))
                if kind == "TOTAL":
                    for c in range(1, ncols + 1):
                        ws.cell(row=r, column=c).font = TOTAL_FONT
            for c in range(1, ncols + 1):
                ws.cell(row=r, column=c).border = BORDER
            r += 1

        self._write_source_cell(ws, r + 1, ncols, sources_text, height=source_height)
        self._autosize(ws, [first_col_width] + [15] * len(years))
        return ws

    def add_cash_flow_sheet(
        self,
        title,
        subtitle,
        rows,
        sources_text,
        sheet_name="Cash Flow Statement",
        first_col_width=60,
        source_height=150,
        unit_suffix=" (£'000)",
        years=None,
    ):
        """
        rows: list of (kind, label, values) - kind in {"SECTION", "DATA", "TOTAL"}.
        years: optional override of self.years for this sheet only - see
            _add_statement_sheet's docstring.
        """
        return self._add_statement_sheet(
            title,
            subtitle,
            rows,
            sources_text,
            sheet_name,
            first_col_width,
            source_height,
            unit_suffix,
            years=years,
        )

    def add_balance_sheet_sheet(
        self,
        title,
        subtitle,
        rows,
        sources_text,
        sheet_name="Balance Sheet",
        first_col_width=60,
        source_height=150,
        unit_suffix=" (£'000)",
        years=None,
    ):
        """Consolidated Statement of Financial Position. Same row/column
        shape as add_cash_flow_sheet - see that docstring for `rows`/`years`."""
        return self._add_statement_sheet(
            title,
            subtitle,
            rows,
            sources_text,
            sheet_name,
            first_col_width,
            source_height,
            unit_suffix,
            years=years,
        )

    def add_income_statement_sheet(
        self,
        title,
        subtitle,
        rows,
        sources_text,
        sheet_name="Profit & Loss",
        first_col_width=60,
        source_height=150,
        unit_suffix=" (£'000)",
        years=None,
    ):
        """Consolidated Statement of Comprehensive Income. Same row/column
        shape as add_cash_flow_sheet - see that docstring for `rows`/`years`."""
        return self._add_statement_sheet(
            title,
            subtitle,
            rows,
            sources_text,
            sheet_name,
            first_col_width,
            source_height,
            unit_suffix,
            years=years,
        )

    def add_equity_changes_sheet(
        self,
        title,
        subtitle,
        headers,
        rows,
        sources_text,
        sheet_name="Statement of Changes in Equity",
        first_col_width=46,
        source_height=150,
        col_width=15,
    ):
        """Consolidated Statement of Changes in Equity.

        Unlike every other statement sheet, this one doesn't fit the
        year-column shape - it's an equity-component-column, movement-row
        roll-forward, read chronologically oldest-to-newest (opening
        balance, this year's movements, closing balance, next year's
        movements, ...), not most-recent-first like the rest of the
        workbook.

        headers: list of equity component names (e.g. "Share capital",
            "Share premium", "Other reserves", "Merger reserve", "Retained
            losses", "Total equity") - becomes the column headers after the
            leading "Movement" label column.
        rows: list of (kind, label, values) - kind in {"DATA", "TOTAL"};
            TOTAL bolds the row (use it for balance-brought/carried-forward
            rows and "Total comprehensive income/(loss) for the year" rows).
            `values` is a tuple/list positionally aligned to `headers`
            (not a year dict) - use None for a component not applicable to
            that row (e.g. Merger reserve before it existed).
        """
        ws = self._next_sheet(sheet_name)
        ws.freeze_panes = "A5"
        ws["A1"] = title
        ws["A1"].font = TITLE_FONT
        ws["A2"] = subtitle
        ws["A2"].font = SUBTITLE_FONT

        ncols = 1 + len(headers)
        col_headers = ["Movement"] + list(headers)
        header_row = 3
        for c, h in enumerate(col_headers, start=1):
            ws.cell(row=header_row, column=c, value=h)
        self._style_header(ws, header_row, ncols)

        r = header_row + 1
        for kind, label, values in rows:
            ws.cell(row=r, column=1, value=label)
            for ci, v in enumerate(values, start=2):
                ws.cell(row=r, column=ci, value=v)
            if kind == "TOTAL":
                for c in range(1, ncols + 1):
                    ws.cell(row=r, column=c).font = TOTAL_FONT
            for c in range(1, ncols + 1):
                ws.cell(row=r, column=c).border = BORDER
            r += 1

        self._write_source_cell(ws, r + 1, ncols, sources_text, height=source_height)
        self._autosize(ws, [first_col_width] + [col_width] * len(headers))
        return ws

    def add_asset_quality_sheet(
        self,
        title,
        subtitle,
        rows,
        sources_text,
        sheet_name="Asset Quality",
        first_col_width=58,
        source_height=150,
        unit_suffix=" (£'000)",
        years=None,
    ):
        """Asset Quality / Credit Risk Disclosures: loan book breakdown by
        product and by IFRS 9 stage (1/2/3), plus derived coverage/NPL
        ratios. Same year-column row/column shape as add_cash_flow_sheet -
        see that docstring for `rows`. Ratio rows may use string values
        (e.g. "13.55%") the same way add_metric_sheet does. Placed right
        after Cash Flow Statement and right before the Pillar 3 sheets.
        years: optional per-sheet override - see _add_statement_sheet."""
        return self._add_statement_sheet(
            title,
            subtitle,
            rows,
            sources_text,
            sheet_name,
            first_col_width,
            source_height,
            unit_suffix,
            years=years,
        )

    def add_rwa_breakdown_sheet(
        self,
        title,
        subtitle,
        rows,
        sources_text,
        sheet_name="RWA Breakdown",
        first_col_width=54,
        source_height=150,
        unit_suffix=" (£'000)",
        years=None,
    ):
        """RWA breakdown by risk category (Pillar 3's UK OV1 template:
        credit risk, counterparty credit risk, securitisation, market risk,
        operational risk). Same year-column row/column shape as
        add_cash_flow_sheet - see that docstring for `rows`. Placed with the
        other Pillar 3 sheets (see PILLAR3_SHEET_NAMES), not with the other
        3 statement sheets, since it's itself a Pillar 3 disclosure.
        years: optional per-sheet override - see _add_statement_sheet."""
        return self._add_statement_sheet(
            title,
            subtitle,
            rows,
            sources_text,
            sheet_name,
            first_col_width,
            source_height,
            unit_suffix,
            years=years,
        )

    def add_km1_sheet(
        self,
        title,
        subtitle,
        rows,
        sources_text,
        sheet_name=KM1_SHEET_NAME,
        first_col_width=72,
        source_height=150,
        unit_suffix="",
        years=None,
    ):
        """The bank's own UK KM1 "Key metrics" template, metrics as rows and
        years as columns, placed immediately BEFORE the 11 individual Pillar 3
        metric sheets so a reader meets the published summary first and the
        single-metric sheets afterwards.

        Same (kind, label, values) row shape as add_cash_flow_sheet - see that
        docstring - but with two deliberate differences:

        `unit_suffix` DEFAULTS TO EMPTY, because KM1 is the one sheet that
        mixes units down a single column: amount rows are £'000 while ratio
        and buffer rows are percentages. Carry the unit on each row instead,
        either in the label (the template's own convention, e.g. "Total
        exposure measure (£'000)") or by writing ratio values as strings
        ("14.5%"), which _add_statement_sheet passes through unchanged.

        TRANSCRIBE THE TEMPLATE AS THE BANK PRINTED IT. KM1 is a prescribed
        template (CRR Article 447 / PRA Disclosure Rules), so fidelity means
        keeping the bank's own row order, its row labels, AND the template's
        row numbers as printed ("UK 7a", "8", "UK 9a", ...) - they are how a
        reader cross-refers to the source document. Keep rows the bank prints
        as blank, "-" or "N/A" rather than dropping them: an omitted row is
        indistinguishable from a row we failed to find. Use SECTION rows for
        the template's own headings ("Additional own funds requirements",
        "Combined buffer requirement", "Leverage ratio", "Liquidity Coverage
        Ratio", "Net Stable Funding Ratio").

        Do NOT reorder rows into this project's preferred sequence, do not
        merge the two sides of a basis break into one row (the 1 Jan 2022
        leverage change and the LCR average-vs-point-in-time split both
        surface here as separate template rows), and do not compute a row the
        bank left blank - a derived figure in a disclosure template reads as
        a disclosure.

        years: optional per-sheet override - see _add_statement_sheet."""
        return self._add_statement_sheet(
            title,
            subtitle,
            rows,
            sources_text,
            sheet_name,
            first_col_width,
            source_height,
            unit_suffix,
            years=years,
        )

    def add_metric_sheet(
        self,
        name,
        unit,
        rows_data,
        sources_text,
        note=None,
        first_col_width=46,
        source_height=150,
        note_height=60,
        years=None,
    ):
        """
        rows_data: list of (label, values) pairs, values a {year: value} dict
            (numbers or strings, e.g. "82.94%" / "Not publicly disclosed").
        years: optional per-sheet override - see _add_statement_sheet.
        """
        years = years if years is not None else self.years
        years = self._trim_trailing_empty_years(
            years, (values for _label, values in rows_data)
        )
        ws = self.wb.create_sheet(title=name[:31])
        ncols = 1 + len(years)
        ws["A1"] = f"{self.bank_name} — {name}"
        ws["A1"].font = TITLE_FONT
        ws["A2"] = unit or ""
        ws["A2"].font = SUBTITLE_FONT

        headers = ["Metric"] + [self.year_label[y] for y in years]
        header_row = 3
        for c, h in enumerate(headers, start=1):
            ws.cell(row=header_row, column=c, value=h)
        self._style_header(ws, header_row, ncols)

        r = header_row + 1
        for label, values in rows_data:
            ws.cell(row=r, column=1, value=label)
            for ci, y in enumerate(years, start=2):
                ws.cell(row=r, column=ci, value=values.get(y))
            for c in range(1, ncols + 1):
                ws.cell(row=r, column=c).border = BORDER
            r += 1

        if note:
            ws.cell(row=r, column=1, value=f"Note: {note}").font = SUBTITLE_FONT
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
            ws.cell(row=r, column=1).alignment = Alignment(
                wrap_text=True, vertical="top"
            )
            ws.row_dimensions[r].height = note_height
            r += 1

        self._write_source_cell(ws, r + 1, ncols, sources_text, height=source_height)
        self._autosize(ws, [first_col_width] + [15] * len(years))
        return ws

    def add_long_form_sheet(
        self,
        name,
        headers,
        rows,
        title=None,
        subtitle=None,
        note=None,
        widths=None,
        hyperlink_cells=None,
        freeze_panes="A5",
    ):
        """Add an auxiliary long-form data sheet.

        This is intended for additional observations that do not fit the
        standard fixed annual columns, such as quarterly or semi-annual Pillar
        3 disclosures.

        ``headers`` is a list of column names and ``rows`` is a list of row
        lists/tuples with the same length.  ``hyperlink_cells`` is an optional
        mapping of ``(data_row_index, column_index)`` to URL, using zero-based
        indexes relative to the supplied data rows and headers.  ``widths`` is
        an optional list of column widths.
        """
        if not headers:
            raise ValueError("Long-form sheet requires at least one header")
        ncols = len(headers)
        for row in rows:
            if len(row) != ncols:
                raise ValueError(
                    f"Long-form row has {len(row)} values; expected {ncols}"
                )

        ws = self.wb.create_sheet(title=name[:31])
        ws["A1"] = title or f"{self.bank_name} — {name}"
        ws["A1"].font = TITLE_FONT
        ws["A2"] = subtitle or ""
        ws["A2"].font = SUBTITLE_FONT

        header_row = 4
        for col, header in enumerate(headers, start=1):
            ws.cell(row=header_row, column=col, value=header)
        self._style_header(ws, header_row, ncols)

        hyperlink_cells = hyperlink_cells or {}
        for row_index, values in enumerate(rows):
            excel_row = header_row + 1 + row_index
            for col_index, value in enumerate(values):
                cell = ws.cell(row=excel_row, column=col_index + 1, value=value)
                cell.border = BORDER
                url = hyperlink_cells.get((row_index, col_index))
                if url:
                    cell.hyperlink = url
                    cell.style = "Hyperlink"

        if note:
            note_row = header_row + 1 + len(rows) + 1
            ws.cell(row=note_row, column=1, value=note)
            ws.merge_cells(
                start_row=note_row,
                start_column=1,
                end_row=note_row,
                end_column=ncols,
            )
            ws.cell(row=note_row, column=1).font = SOURCE_FONT
            ws.cell(row=note_row, column=1).alignment = Alignment(
                wrap_text=True, vertical="top", horizontal="left"
            )
            ws.row_dimensions[note_row].height = 60

        ws.freeze_panes = freeze_panes
        self._autosize(ws, widths or [18] * ncols)
        return ws

    def add_wide_interim_sheet(
        self,
        name,
        headers=None,
        rows=None,
        title=None,
        subtitle=None,
        note=None,
        widths=None,
        hyperlink_cells=None,
        freeze_panes="D5",
    ):
        """Add a wide interim Pillar 3 matrix from standard observation rows.

        ``rows`` uses the same eight-field records as ``add_long_form_sheet``:
        period, disclosure type, metric, value, unit, basis, source document,
        and page/table.  Periods become columns and metrics become rows.  Each
        populated value cell links to its source document; a compact source
        register below the matrix preserves disclosure type and page/table
        references.  ``hyperlink_cells`` may map the zero-based input row and
        source-document column (normally ``(row_index, 6)``) to a URL.  This
        keeps a readable source-document label in the workbook while using the
        official URL as the hyperlink target.
        """
        # Accept the existing long-form call shape as well as the simpler
        # ``rows=...`` form, so bank builders can be migrated mechanically.
        if rows is None:
            rows = headers
        if not rows:
            raise ValueError("Wide interim sheet requires at least one row")
        if any(len(row) != 8 for row in rows):
            raise ValueError("Wide interim rows must contain exactly 8 fields")

        periods = []
        metrics = []
        metric_meta = {}
        values = {}
        links = {}
        source_rows = []
        source_seen = set()
        hyperlink_cells = hyperlink_cells or {}
        for row_index, row in enumerate(rows):
            period, disclosure, metric, value, unit, basis, source, page = row
            if period not in periods:
                periods.append(period)
            if metric not in metrics:
                metrics.append(metric)
                metric_meta[metric] = (unit, basis)
            values[(metric, period)] = value
            # Only ever hyperlink a cell to a genuine URL. This used to fall
            # back to the plain source-document label text (e.g. "Q3 2022
            # Pillar 3 Supplement") whenever a build script didn't pass an
            # explicit URL for that row - openpyxl writes that text in
            # verbatim as `cell.hyperlink`, producing an invalid link Excel
            # either errors on or mistreats as a relative path, and it LOOKS
            # like a real citation (styled + clickable) even though it isn't
            # one. Leaving the cell as plain, unstyled text when no real URL
            # is available makes the gap visible instead of silently faking
            # a working citation - see IN-031.
            source_url = hyperlink_cells.get((row_index, 6))
            if source_url and not str(source_url).startswith(("http://", "https://")):
                source_url = None
            links[(metric, period)] = source_url
            source_key = (period, disclosure, source, page)
            if source_key not in source_seen:
                source_seen.add(source_key)
                source_rows.append((*source_key, source_url))

        ws = self.wb.create_sheet(title=name[:31])
        ws["A1"] = title or f"{self.bank_name} — {name}"
        ws["A1"].font = TITLE_FONT
        ws["A2"] = subtitle or ""
        ws["A2"].font = SUBTITLE_FONT

        headers = ["Metric", "Unit", "Basis"] + periods
        header_row = 4
        for col, header in enumerate(headers, start=1):
            ws.cell(row=header_row, column=col, value=header)
        self._style_header(ws, header_row, len(headers))

        for row_index, metric in enumerate(metrics):
            excel_row = header_row + 1 + row_index
            unit, basis = metric_meta[metric]
            row_values = [metric, unit, basis]
            for col_index, value in enumerate(row_values, start=1):
                cell = ws.cell(row=excel_row, column=col_index, value=value)
                cell.border = BORDER
            for period_index, period in enumerate(periods, start=4):
                value = values.get((metric, period))
                cell = ws.cell(row=excel_row, column=period_index, value=value)
                cell.border = BORDER
                source = links.get((metric, period))
                if source and value not in (None, ""):
                    cell.hyperlink = source
                    cell.style = "Hyperlink"

        source_header_row = header_row + len(metrics) + 3
        ws.cell(row=source_header_row, column=1, value="Source register")
        ws.cell(row=source_header_row, column=1).font = SUBTITLE_FONT
        source_headers = [
            "Period",
            "Disclosure type",
            "Source document",
            "Page / table",
        ]
        for col, header in enumerate(source_headers, start=1):
            ws.cell(row=source_header_row + 1, column=col, value=header)
        self._style_header(ws, source_header_row + 1, len(source_headers))
        for source_index, source_row in enumerate(source_rows):
            excel_row = source_header_row + 2 + source_index
            for col_index, value in enumerate(source_row[:4], start=1):
                cell = ws.cell(row=excel_row, column=col_index, value=value)
                cell.border = BORDER
                if col_index == 3 and source_row[4]:
                    cell.hyperlink = source_row[4]
                    cell.style = "Hyperlink"

        note_row = source_header_row + 2 + len(source_rows) + 1
        if note:
            ws.cell(row=note_row, column=1, value=f"Note: {note}")
            ws.merge_cells(
                start_row=note_row,
                start_column=1,
                end_row=note_row,
                end_column=len(headers),
            )
            ws.cell(row=note_row, column=1).font = SOURCE_FONT
            ws.cell(row=note_row, column=1).alignment = Alignment(
                wrap_text=True, vertical="top"
            )
            ws.row_dimensions[note_row].height = 60

        ws.freeze_panes = freeze_panes
        self._autosize(ws, [32, 18, 34] + [15] * len(periods))
        return ws

    def add_overview_sheet(
        self,
        cash_flow_totals,
        cash_flow_unit,
        ratios,
        note=None,
        balance_sheet_totals=None,
        balance_sheet_unit=None,
        income_statement_totals=None,
        income_statement_unit=None,
        equity_changes_totals=None,
        equity_changes_unit=None,
        years=None,
    ):
        """
        Inserts an "Overview" sheet as the first tab: summary tables of
        headline Balance Sheet / Profit & Loss / Statement of Changes in
        Equity / cash flow totals and headline Pillar 3 ratios across all
        years, plus a clustered column chart for each money block and a
        line chart for the ratios - all rendered oldest-year-first left to
        right regardless of the most-recent-first column order used
        everywhere else. Block order, top to bottom: Balance Sheet,
        Profit & Loss, Statement of Changes in Equity, Cash Flow, Ratios.

        balance_sheet_totals / income_statement_totals: list of
            (label, {year: number}) - e.g. Total assets/Loans and advances/
            Customer deposits/Total equity for the balance sheet; Revenue/
            Total operating expense/Profit for the year for the income
            statement. Omit or pass `None`/`[]` to skip that block entirely
            (e.g. for a workbook built before these sheets existed).
        balance_sheet_unit / income_statement_unit: unit string for that
            block's table heading and chart axis, e.g. "£'000". Ignored if
            the corresponding totals list is empty.
        equity_changes_totals: list of (label, {year: number}) - a per-year
            bridge summary of the Statement of Changes in Equity sheet
            (which is itself chronological, not year-columned - see
            add_equity_changes_sheet), e.g. Opening equity / Total
            comprehensive income(/loss) for the year / Other equity
            movements, net / Closing equity. Omit or pass `None`/`[]` to
            skip.
        equity_changes_unit: unit string for that block, e.g. "£'000".
            Ignored if equity_changes_totals is empty.
        cash_flow_totals: list of (label, {year: number}) - e.g. net cash
            from operating/investing/financing activities, cash at year end.
            Pass an empty list ([]) for a Pillar-3-only workbook (e.g. an
            entity taking the FRS 101/102 cash-flow-statement exemption) -
            the cash-flow table/chart are then omitted entirely; see the
            "Cash Flow Statement" sheet for the exemption note in that case.
        cash_flow_unit: unit string for the table heading and chart axis,
            e.g. "£m" or "£'000". Ignored if cash_flow_totals is empty.
        ratios: list of (label, {year: value}) - values as displayed
            (e.g. "14.5%" or "Not publicly disclosed"); values that aren't
            a plain percentage are parsed as None and simply leave a gap
            in the chart rather than being plotted. Pass an empty list ([])
            for a bank with no disclosed Pillar 3 ratios at all - the ratios
            table/chart are then omitted entirely (fixed 2026-08-28; earlier
            build scripts worked around this by passing empty-{}-valued
            entries instead, e.g. build_caf_bank.py/build_brown_shipley.py -
            that workaround still works too, no need to change them).
        """
        years = years if years is not None else self.years
        ws = self.wb.create_sheet(title="Overview", index=0)
        ncols = 1 + len(years)
        headers = ["Line item"] + [self.year_label[y] for y in years]

        ws["A1"] = f"{self.bank_name} — Overview"
        ws["A1"].font = TITLE_FONT
        ws["A2"] = (
            "Headline balance sheet, profit & loss, cash flow and Pillar 3 figures across all years "
            "covered in this workbook - how the bank is using its money and the risk it is taking with it"
        )
        ws["A2"].font = SUBTITLE_FONT

        # -- money blocks: Balance Sheet, Profit & Loss, Equity Changes, Cash Flow (bar charts) --
        money_blocks = [
            (
                "bs",
                "Balance Sheet Summary",
                balance_sheet_totals or [],
                balance_sheet_unit,
            ),
            (
                "is",
                "Profit & Loss Summary",
                income_statement_totals or [],
                income_statement_unit,
            ),
            (
                "eq",
                "Statement of Changes in Equity Summary",
                equity_changes_totals or [],
                equity_changes_unit,
            ),
            ("cf", "Cash Flow Summary", cash_flow_totals or [], cash_flow_unit),
        ]
        not_applicable_text = {
            "bs": "Balance Sheet Summary: not applicable — this workbook was built before the Balance Sheet sheet existed.",
            "is": "Profit & Loss Summary: not applicable — this workbook was built before the Profit & Loss sheet existed.",
            "eq": "Statement of Changes in Equity Summary: not applicable — this workbook was built before the Statement of Changes in Equity sheet existed.",
            "cf": "Cash Flow Summary: not applicable — this entity does not publish a cash flow statement (see the Cash Flow Statement sheet for the exemption basis).",
        }

        row = 4
        block_table_info = {}  # key -> (header_row, data_end) or None
        for key, title, totals, unit in money_blocks:
            if totals:
                ws.cell(
                    row=row, column=1, value=f"{title} ({unit})"
                ).font = SECTION_FONT
                row += 1
                header_row = row
                for c, h in enumerate(headers, start=1):
                    ws.cell(row=row, column=c, value=h)
                self._style_header(ws, row, ncols)
                row += 1
                for label, values in totals:
                    ws.cell(row=row, column=1, value=label)
                    for ci, y in enumerate(years, start=2):
                        ws.cell(row=row, column=ci, value=values.get(y))
                    for c in range(1, ncols + 1):
                        ws.cell(row=row, column=c).font = TOTAL_FONT
                        ws.cell(row=row, column=c).border = BORDER
                    row += 1
                block_table_info[key] = (header_row, row - 1)
            else:
                ws.cell(
                    row=row, column=1, value=not_applicable_text[key]
                ).font = SECTION_FONT
                row += 1
                block_table_info[key] = None
            row += 1

        if ratios:
            ws.cell(row=row, column=1, value="Pillar 3 Key Metrics").font = SECTION_FONT
            row += 1
            ratio_header_row = row
            for c, h in enumerate(headers, start=1):
                ws.cell(row=row, column=c, value=h)
            self._style_header(ws, row, ncols)
            row += 1
            for label, values in ratios:
                ws.cell(row=row, column=1, value=label)
                for ci, y in enumerate(years, start=2):
                    ws.cell(row=row, column=ci, value=values.get(y))
                for c in range(1, ncols + 1):
                    ws.cell(row=row, column=c).border = BORDER
                row += 1
            ratio_data_end = row - 1
        else:
            ws.cell(
                row=row,
                column=1,
                value="Pillar 3 Key Metrics: not applicable — no ratios are disclosed for this "
                "entity (see the individual Pillar 3 metric sheets).",
            ).font = SECTION_FONT
            row += 1
            ratio_header_row = None

        if note:
            row += 1
            ws.cell(row=row, column=1, value=f"Note: {note}").font = SUBTITLE_FONT
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncols)
            ws.cell(row=row, column=1).alignment = Alignment(
                wrap_text=True, vertical="top"
            )
            ws.row_dimensions[row].height = 45
            row += 1

        self._autosize(ws, [42] + [15] * len(years))

        # -- hidden numeric staging areas, oldest-year-first, for the charts --
        # (built as separate chronological blocks rather than via a reversed
        # category axis, which flips the value axis to the wrong side in Excel)
        chrono_years = list(reversed(years))
        n_years = len(years)

        stage_col = ncols + 3
        block_stage_info = {}  # key -> (stage_header_row, stage_end) or None
        for key, title, totals, unit in money_blocks:
            info = block_table_info[key]
            if totals and info:
                header_row, _ = info
                for ci, y in enumerate(chrono_years, start=1):
                    ws.cell(
                        row=header_row, column=stage_col + ci, value=self.year_label[y]
                    )
                r = header_row + 1
                for label, values in totals:
                    ws.cell(row=r, column=stage_col, value=label)
                    for ci, y in enumerate(chrono_years, start=1):
                        ws.cell(row=r, column=stage_col + ci, value=values.get(y))
                    r += 1
                block_stage_info[key] = (header_row, r - 1, stage_col)
                stage_col += ncols + 2
            else:
                block_stage_info[key] = None

        ratio_stage_col = stage_col
        if ratios:
            ratio_stage_header_row = ratio_header_row
            for prev_key in ("bs", "is", "eq", "cf"):
                if block_stage_info[prev_key]:
                    ratio_stage_header_row = block_stage_info[prev_key][0]
                    break
            for ci, y in enumerate(chrono_years, start=1):
                ws.cell(
                    row=ratio_stage_header_row,
                    column=ratio_stage_col + ci,
                    value=self.year_label[y],
                )
            r = ratio_stage_header_row + 1
            for label, values in ratios:
                ws.cell(row=r, column=ratio_stage_col, value=label)
                for ci, y in enumerate(chrono_years, start=1):
                    ws.cell(
                        row=r,
                        column=ratio_stage_col + ci,
                        value=_parse_percent(values.get(y)),
                    )
                r += 1
            ratio_stage_end = r - 1
            last_col = ratio_stage_col + ncols
        else:
            last_col = stage_col

        for c in range(ncols + 3, last_col):
            ws.column_dimensions[get_column_letter(c)].hidden = True

        # -- charts --
        chart_row = row + 2
        bar_titles = {
            "bs": "Balance Sheet Summary",
            "is": "Profit & Loss Summary",
            "eq": "Statement of Changes in Equity Summary",
            "cf": "Cash Flow Summary",
        }
        for key, title, totals, unit in money_blocks:
            stage_info = block_stage_info[key]
            if not stage_info:
                continue
            stage_header_row, stage_end, block_stage_col = stage_info
            bar = BarChart()
            bar.type = "col"
            bar.grouping = "clustered"
            bar.title = f"{self.bank_name} — {bar_titles[key]} by Year"
            bar.y_axis.title = unit
            bar.height = 13
            bar.width = 24
            data = Reference(
                ws,
                min_col=block_stage_col,
                max_col=block_stage_col + n_years,
                min_row=stage_header_row + 1,
                max_row=stage_end,
            )
            bar.add_data(data, titles_from_data=True, from_rows=True)
            cats = Reference(
                ws,
                min_col=block_stage_col + 1,
                max_col=block_stage_col + n_years,
                min_row=stage_header_row,
                max_row=stage_header_row,
            )
            bar.set_categories(cats)
            bar.visible_cells_only = (
                False  # source data lives in hidden staging columns
            )
            self._style_chart_axes(bar)
            bar.legend.position = "b"
            self._reserve_bottom_legend(bar, legend_h=0.16, plot_h=0.55, plot_top=0.14)
            ws.add_chart(bar, f"A{chart_row}")
            chart_row += 22

        if not ratios:
            return ws

        line = LineChart()
        line.title = f"{self.bank_name} — Pillar 3 Key Metrics by Year"
        line.y_axis.title = "%"
        line.height = 16
        line.width = 24
        ldata = Reference(
            ws,
            min_col=ratio_stage_col,
            max_col=ratio_stage_col + n_years,
            min_row=ratio_stage_header_row + 1,
            max_row=ratio_stage_end,
        )
        line.add_data(ldata, titles_from_data=True, from_rows=True)
        lcats = Reference(
            ws,
            min_col=ratio_stage_col + 1,
            max_col=ratio_stage_col + n_years,
            min_row=ratio_stage_header_row,
            max_row=ratio_stage_header_row,
        )
        line.set_categories(lcats)
        line.visible_cells_only = False  # source data lives in hidden staging columns
        self._style_chart_axes(line)
        line.legend.position = "b"
        self._reserve_bottom_legend(line, legend_h=0.22, plot_h=0.50, plot_top=0.12)
        for s in line.series:
            s.smooth = False
        ws.add_chart(line, f"A{chart_row}")

        return ws

    def add_not_disclosed_metric_sheets(self, names, sources_text, per_note=None, years=None,
                                        source_height=150, statements=None):
        """Convenience for the common "no Pillar 3 doc exists" case - fills a
        list of metric sheets with a single "Not publicly disclosed" row each.
        per_note: optional {name: note_text} for sheet-specific notes.
        years: optional per-sheet override - see _add_statement_sheet.
        source_height: row height of the merged source-citation cell. The
        default matches _write_source_cell's own default; raise it when
        sources_text is long, since that cell has a fixed height and simply
        clips any citation text that overflows it.
        statements: optional {name: text} or {name: {year: text}} replacing the
        bare "Not publicly disclosed" with a statement that says WHICH outcome
        it is (GA-020): "Not published – <evidence>", "Not applicable –
        <reason>", "Unreached today – <what was tried>" or "Not published yet
        – <when due>". Omitted names/years keep the old bare text, so existing
        callers are unchanged."""
        per_note = per_note or {}
        statements = statements or {}
        years = years if years is not None else self.years
        for name in names:
            st = statements.get(name)
            cells = {}
            for y in years:
                if isinstance(st, dict):
                    cells[y] = st.get(y, "Not publicly disclosed")
                elif st:
                    cells[y] = st
                else:
                    cells[y] = "Not publicly disclosed"
            self.add_metric_sheet(
                name,
                None,
                [(name, cells)],
                sources_text,
                note=per_note.get(name),
                years=years,
                source_height=source_height,
            )

    def save(self, path):
        self._assert_no_orphan_columns()
        self.wb.save(path)
        print("Saved workbook with sheets:", self.wb.sheetnames)
        return path
