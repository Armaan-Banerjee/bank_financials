"""Renders a single bank's real .xlsx workbook (from banks/) as a
self-contained, tabbed static HTML page - one tab per sheet, laid out as a
plain HTML table with merged cells, bold rows, and header fills carried over
from the workbook's own openpyxl styling. No JS spreadsheet library, no
server: this is meant to be embedded via <iframe> in a bank-<slug>.html
drilldown page (imported and called from build_deliverable.py), or opened
standalone.

Deliberately simple - this is a faithful-enough READ view of what's already
in the workbook, not an interactive spreadsheet. Formulas aren't evaluated
here since data_only=True reads openpyxl's last-cached value, matching what
Excel itself last computed and saved.
"""

import html as htmlmod

import openpyxl

WORKBOOK_VIEWER_CSS = """
* { box-sizing: border-box; }
body { margin:0; font-family:-apple-system,BlinkMacSystemFont,'Segoe UI',Helvetica,Arial,sans-serif; font-size:12px; color:#2a2a28; background:#fff; }
.wb-tabs { display:flex; flex-wrap:wrap; gap:2px; padding:6px 6px 0; background:#f4f2ec; border-bottom:1px solid #ddd8ca; position:sticky; top:0; z-index:1; }
.wb-tab { border:1px solid #ddd8ca; border-bottom:none; background:#ece9e0; color:#5b5647; padding:5px 11px; font-size:11px; cursor:pointer; border-radius:3px 3px 0 0; }
.wb-tab.active { background:#fff; color:#2a2a28; font-weight:600; }
.wb-panel { display:none; overflow:auto; max-height:560px; }
.wb-panel.active { display:block; }
.wb-table { border-collapse:collapse; font-size:11.5px; white-space:nowrap; }
.wb-table td { padding:4px 9px; border:1px solid #ece9e0; vertical-align:middle; }
"""

WORKBOOK_VIEWER_JS = """
document.querySelectorAll('.wb-tab').forEach(tab => {
  tab.addEventListener('click', () => {
    document.querySelectorAll('.wb-tab').forEach(t => t.classList.remove('active'));
    document.querySelectorAll('.wb-panel').forEach(p => p.classList.remove('active'));
    tab.classList.add('active');
    document.querySelector(`.wb-panel[data-idx="${tab.dataset.idx}"]`).classList.add('active');
  });
});
"""


def _cell_style(cell):
    style = []
    if cell.font and cell.font.bold:
        style.append("font-weight:600")
    if cell.font and cell.font.color and cell.font.color.rgb and isinstance(cell.font.color.rgb, str):
        rgb = cell.font.color.rgb
        if len(rgb) == 8 and rgb != "00000000":
            style.append(f"color:#{rgb[-6:]}")
    if cell.fill and cell.fill.patternType == "solid" and cell.fill.fgColor and isinstance(cell.fill.fgColor.rgb, str):
        rgb = cell.fill.fgColor.rgb
        if len(rgb) == 8 and rgb not in ("00000000",):
            style.append(f"background:#{rgb[-6:]}")
    align = cell.alignment.horizontal if cell.alignment else None
    if align:
        style.append(f"text-align:{align}")
    if cell.alignment and cell.alignment.wrap_text:
        style.append("white-space:normal")
    return ";".join(style)


def _sheet_to_table_html(ws):
    merged_starts = {}
    skip = set()
    for rng in ws.merged_cells.ranges:
        merged_starts[(rng.min_row, rng.min_col)] = (rng.max_row - rng.min_row + 1, rng.max_col - rng.min_col + 1)
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                if (r, c) != (rng.min_row, rng.min_col):
                    skip.add((r, c))

    rows_html = []
    max_row = ws.max_row or 0
    max_col = ws.max_column or 0
    for r in range(1, max_row + 1):
        cells_html = []
        for c in range(1, max_col + 1):
            if (r, c) in skip:
                continue
            cell = ws.cell(row=r, column=c)
            span = merged_starts.get((r, c))
            rowspan = f' rowspan="{span[0]}"' if span and span[0] > 1 else ""
            colspan = f' colspan="{span[1]}"' if span and span[1] > 1 else ""
            value = cell.value
            text = "" if value is None else str(value)
            text = htmlmod.escape(text).replace("\n", "<br>")
            style = _cell_style(cell)
            style_attr = f' style="{style}"' if style else ""
            cells_html.append(f"<td{rowspan}{colspan}{style_attr}>{text}</td>")
        rows_html.append("<tr>" + "".join(cells_html) + "</tr>")
    return '<table class="wb-table">' + "".join(rows_html) + "</table>"


def render_workbook_viewer(xlsx_path, title):
    """Return a standalone HTML document with one tab per sheet in
    xlsx_path, in the workbook's own sheet order."""
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    sheets = [(ws.title, _sheet_to_table_html(ws)) for ws in wb.worksheets]

    tabs = "".join(
        f'<button class="wb-tab{" active" if i == 0 else ""}" data-idx="{i}">{htmlmod.escape(name)}</button>'
        for i, (name, _) in enumerate(sheets)
    )
    panels = "".join(
        f'<div class="wb-panel{" active" if i == 0 else ""}" data-idx="{i}">{table}</div>'
        for i, (name, table) in enumerate(sheets)
    )
    return f"""<!doctype html>
<html lang="en">
<head>
<meta charset="utf-8">
<title>{htmlmod.escape(title)}</title>
<meta name="viewport" content="width=device-width, initial-scale=1">
<style>{WORKBOOK_VIEWER_CSS}</style>
</head>
<body>
<div class="wb-tabs">{tabs}</div>
{panels}
<script>{WORKBOOK_VIEWER_JS}</script>
</body>
</html>
"""
