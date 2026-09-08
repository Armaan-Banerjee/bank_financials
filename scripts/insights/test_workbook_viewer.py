"""Focused tests for the static workbook HTML renderer."""

import sys
import tempfile
import unittest
from pathlib import Path

import openpyxl


SCRIPTS = Path(__file__).resolve().parent
sys.path.insert(0, str(SCRIPTS))
from build_workbook_viewer import render_workbook_viewer


class WorkbookViewerTests(unittest.TestCase):
    def test_hidden_chart_staging_columns_are_not_rendered(self):
        """Deep-history workbooks must not expose hidden helper columns."""
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Overview"
        ws["A1"] = "Visible workbook content"
        ws["B1"] = "FY1973"
        ws["C1"] = "chart-staging-only"
        ws.column_dimensions["C"].hidden = True

        with tempfile.TemporaryDirectory(prefix="workbook_viewer_test_") as tmp:
            path = Path(tmp) / "fixture.xlsx"
            wb.save(path)
            rendered = render_workbook_viewer(path, "Fixture")

        self.assertIn("Visible workbook content", rendered)
        self.assertIn("FY1973", rendered)
        self.assertNotIn("chart-staging-only", rendered)


if __name__ == "__main__":
    unittest.main(verbosity=2)
